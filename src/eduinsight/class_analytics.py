"""Course-isolated class analytics and transcripts derived from existing stores.

Enrollment is defined by chat_sessions, as in courses.py. Knowledge facts have
student ownership only; radar responses explicitly disclose that scope. Hours
are estimates using analytics_extra's activity spans, never invented durations.
"""

from __future__ import annotations

import io
import os
from collections import Counter, defaultdict
from dataclasses import asdict
from datetime import datetime, timedelta
from pathlib import Path
from threading import Lock
from urllib.parse import quote
from xml.sax.saxutils import escape

from fastapi import APIRouter, Depends, HTTPException, Response

from . import analytics_extra as extra
from . import assignments, sessions
from .analytics_extra import _activity_spans, _knowledge, _messages, _timestamp
from .auth import TEACHER_COURSE_MAP, AuthUser, get_current_user, require_role
from .courses import COURSE_CATALOG

router = APIRouter(tags=["class-analytics"])
_WEEKDAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
_FONT_LOCK = Lock()


def _course_catalog() -> dict[str, str]:
    """Read dynamic catalog entries when available, without acquiring a write lock."""
    catalog = dict(COURSE_CATALOG)
    with sessions._db() as conn:
        if conn.execute("SELECT 1 FROM sqlite_master WHERE type = 'table' "
                        "AND name = 'course_catalog'").fetchone():
            catalog.update({row["id"]: row["name"] for row in conn.execute(
                "SELECT id, name FROM course_catalog",
            )})
    return catalog


def _own_course_ids(user: AuthUser) -> set[str]:
    with sessions._db() as conn:
        enrolled = {row[0] for row in conn.execute(
            "SELECT DISTINCT course_id FROM chat_sessions WHERE moodle_user_id = ?",
            (user.moodle_user_id,),
        )}
    return enrolled & set(_course_catalog())


def _check_course_access(user: AuthUser, course_id: str) -> None:
    # Analytics teachers remain scoped by TEACHER_COURSE_MAP even if the editable
    # catalog assigns additional courses. Authorization performs only reads.
    if user.role == "teacher" and TEACHER_COURSE_MAP.get(user.username) != course_id:
        raise HTTPException(403, "Teacher course access denied")
    if course_id not in _course_catalog():
        raise HTTPException(404, "Course not found")
    if user.role == "student" and course_id not in _own_course_ids(user):
        raise HTTPException(403, "Course access denied")


def _course_sessions(course_id: str) -> list[dict]:
    with sessions._db() as conn:
        return [dict(row) for row in conn.execute(
            "SELECT * FROM chat_sessions WHERE course_id = ? ORDER BY session_id", (course_id,),
        )]


def _enrolled(course_id: str) -> list[int]:
    return sorted({row["moodle_user_id"] for row in _course_sessions(course_id)})


def _student_access(user: AuthUser, course_id: str, student_id: int) -> None:
    if user.role == "student" and user.moodle_user_id != student_id:
        raise HTTPException(403, "Students may only view their own data")
    _check_course_access(user, course_id)
    if student_id not in _enrolled(course_id):
        raise HTTPException(403, "Student is not enrolled in this course")


def _weekly(rows: list[dict], messages: list[dict]) -> list[dict]:
    """Shared eight-week estimator, with Taipei Monday boundaries."""
    now = extra._now()
    monday = now.replace(hour=0, minute=0, second=0, microsecond=0)
    start = monday - timedelta(days=now.weekday(), weeks=7)
    spans = _activity_spans(rows, messages)
    result = []
    for index in range(8):
        lower = start + timedelta(weeks=index)
        upper = min(lower + timedelta(weeks=1), now).timestamp()
        seconds = sum(max(0, min(end, upper) - max(begin, lower.timestamp()))
                      for begin, end in spans)
        result.append({"week_start": lower.date().isoformat(), "hours": round(seconds / 3600, 4)})
    return result


def _activity_stamps(rows: list[dict], messages: list[dict]) -> list[float]:
    represented = {message["session_id"] for message in messages}
    values = [message["created_at"] for message in messages]
    values.extend(row["last_msg_at"] for row in rows
                  if row["session_id"] not in represented
                  and row["last_msg_at"] > row["created_at"])
    now = extra._now().timestamp()
    # Exclude malformed/out-of-range timestamps before datetime conversion.
    return [ts for value in values if (ts := _timestamp(value)) is not None and 0 <= ts <= now]


def _class_activity(course_id: str) -> list[tuple[int, list[dict], list[dict]]]:
    grouped = defaultdict(list)
    for row in _course_sessions(course_id):
        grouped[row["moodle_user_id"]].append(row)
    memory = extra.get_memory() if grouped else None
    return [(sid, rows, _messages(memory, sid, rows, scoped=True))
            for sid, rows in sorted(grouped.items())]


@router.get("/class-analytics/{course_id}/weekly-hours")
def weekly_hours(course_id: str, user: AuthUser = Depends(require_role("teacher", "admin"))) -> dict:
    _check_course_access(user, course_id)
    activity = _class_activity(course_id)
    weeks = _weekly([], [])
    totals = [0.0] * 8
    for _, rows, messages in activity:
        for index, week in enumerate(_weekly(rows, messages)):
            totals[index] += week["hours"]
    return {"course_id": course_id, "student_count": len(activity), "estimated": True,
            "timezone": "Asia/Taipei", "weeks": [
                {"week_start": week["week_start"], "total_hours": round(total, 4),
                 "average_hours": round(total / len(activity), 4) if activity else 0.0}
                for week, total in zip(weeks, totals)]}


def _heatmap(course_id: str) -> dict:
    counts = [[0] * 24 for _ in range(7)]
    for _, rows, messages in _class_activity(course_id):
        for ts in _activity_stamps(rows, messages):
            when = datetime.fromtimestamp(ts, extra._TAIPEI)
            counts[when.weekday()][when.hour] += 1
    return {"course_id": course_id, "timezone": "Asia/Taipei",
            "weekday_order": _WEEKDAYS, "counts": counts}


@router.get("/class-analytics/{course_id}/focus-heatmap")
def focus_heatmap(course_id: str, user: AuthUser = Depends(require_role("teacher", "admin"))) -> dict:
    _check_course_access(user, course_id)
    return _heatmap(course_id)


@router.get("/class-analytics/{course_id}/golden-hours")
def golden_hours(course_id: str, user: AuthUser = Depends(require_role("teacher", "admin"))) -> dict:
    _check_course_access(user, course_id)
    counts = _heatmap(course_id)["counts"]
    slots = [(counts[day][hour], day, hour) for day in range(7) for hour in range(24)
             if counts[day][hour] > 0]
    slots.sort(key=lambda slot: (-slot[0], slot[1], slot[2]))
    return {"course_id": course_id, "timezone": "Asia/Taipei", "slots": [
        {"weekday": day, "weekday_name": _WEEKDAYS[day], "hour": hour, "count": count}
        for count, day, hour in slots[:3]]}


@router.get("/class-analytics/{course_id}/skills")
def skills(course_id: str, user: AuthUser = Depends(require_role("teacher", "admin"))) -> dict:
    _check_course_access(user, course_id)
    students = _enrolled(course_id)
    topics = defaultdict(Counter)
    memory = extra.get_memory() if students else None
    for sid in students:
        for skill in _knowledge(memory, sid):
            topics[skill["topic"]].update({key: skill[key]
                                         for key in ("mastered", "struggling", "unknown")})
    radar = []
    for topic, counts in sorted(topics.items()):
        evidence = counts["mastered"] + counts["struggling"]
        radar.append({"topic": topic, **counts,
                      "score": round(100 * counts["mastered"] / evidence, 2) if evidence else None})
    return {"course_id": course_id, "student_count": len(students), "anonymised": True,
            "knowledge_scope": "student", "skills": radar}


def _assignment_data(course_id: str) -> tuple[list[dict], dict[tuple[str, int], dict]]:
    enrolled = set(_enrolled(course_id))
    with assignments._db() as conn:
        work = [dict(row) for row in conn.execute(
            "SELECT * FROM assignments WHERE course_id = ? ORDER BY due_at, id", (course_id,),
        )]
        submitted = {(row["assignment_id"], row["student_id"]): dict(row) for row in conn.execute(
            "SELECT s.* FROM assignment_submissions s JOIN assignments a ON a.id = s.assignment_id "
            "WHERE a.course_id = ?", (course_id,),
        ) if row["student_id"] in enrolled}
    return work, submitted


@router.get("/class-analytics/{course_id}/assignment-completion")
def assignment_completion(
    course_id: str, user: AuthUser = Depends(require_role("teacher", "admin")),
) -> dict:
    _check_course_access(user, course_id)
    enrolled = _enrolled(course_id)
    work, submissions = _assignment_data(course_id)
    result = []
    for item in work:
        rows = [submissions.get((item["id"], sid), {}) for sid in enrolled]
        submitted = sum(row.get("state") in ("submitted", "graded") for row in rows)
        result.append({"assignment_id": item["id"], "title": item["title"],
                       "submitted": submitted,
                       "graded": sum(row.get("state") == "graded" for row in rows),
                       "late": sum(bool(row.get("late")) for row in rows
                                   if row.get("state") in ("submitted", "graded")),
                       "missing": len(enrolled) - submitted})
    total = len(work) * len(enrolled)
    return {"course_id": course_id, "student_count": len(enrolled), "assignments": result,
            "overall_completion_rate": round(100 * sum(row["submitted"] for row in result) / total, 2)
            if total else 0.0}


def _display_name(student_id: int) -> str:
    with sessions._db() as conn:
        row = conn.execute("SELECT display_name FROM users WHERE moodle_user_id = ? "
                           "AND role = 'student' ORDER BY id LIMIT 1", (student_id,)).fetchone()
    return row[0] if row else str(student_id)


def _grade_detail(course_id: str, student_id: int, *, knowledge: bool = True) -> dict:
    manager = extra.get_grade_manager()
    records = manager.get_scores(course_id, student_id=student_id)
    categories = []
    weighted_total = None
    try:
        summary = manager.student_summary(course_id, student_id)
        categories = [asdict(category) for category in summary.categories]
        for category in categories:
            if not category["items"]:
                category["average_pct"] = None
        weighted_total = summary.weighted_total if records else None
    except KeyError:
        pass
    indexed = list(enumerate(records))
    indexed.sort(key=lambda pair: (
        _timestamp(getattr(pair[1], "created_at", None)) or float("inf"), pair[0],
    ))
    trend = [{"item": row.item, "score": row.score,
              "at": getattr(row, "created_at", None) or None} for _, row in indexed]
    return {"moodle_user_id": student_id, "display_name": _display_name(student_id),
            "weighted_total": weighted_total, "category_scores": categories,
            "knowledge_scope": "student",
            "knowledge_mastery": _knowledge(extra.get_memory(), student_id) if knowledge else [],
            "grade_trend": trend}


@router.get("/class-analytics/{course_id}/grades-detail")
def grades_detail(course_id: str, user: AuthUser = Depends(require_role("teacher", "admin"))) -> dict:
    _check_course_access(user, course_id)
    return {"course_id": course_id, "students": [_grade_detail(course_id, sid)
                                               for sid in _enrolled(course_id)]}


@router.get("/class-analytics/{course_id}/student/{moodle_user_id}/feedback")
def feedback(course_id: str, moodle_user_id: int,
             user: AuthUser = Depends(get_current_user)) -> dict:
    _student_access(user, course_id, moodle_user_id)
    work, submissions = _assignment_data(course_id)
    items = []
    for item in work:
        row = submissions.get((item["id"], moodle_user_id))
        if row and row["state"] == "graded":
            items.append({"assignment_id": item["id"], "item": item["title"],
                          "score": row["score"], "max_score": item["max_score"],
                          "feedback": row["feedback"], "at": row["graded_at"]})
    items.sort(key=lambda item: (item["at"] or float("inf"), item["assignment_id"]))
    return {"course_id": course_id, "moodle_user_id": moodle_user_id, "feedback": items}


@router.get("/class-analytics/{course_id}/export.xlsx")
def export_excel(course_id: str, user: AuthUser = Depends(require_role("teacher", "admin"))) -> Response:
    _check_course_access(user, course_id)
    from openpyxl import Workbook

    workbook = Workbook()
    overview = workbook.active
    overview.title = "Course"
    students = _enrolled(course_id)
    overview.append(["course_id", "course_name", "student_count"])
    overview.append([course_id, _course_catalog()[course_id], len(students)])
    grades = workbook.create_sheet("Grades")
    grades.append(["student_id", "display_name", "weighted_total", "category", "item", "score", "total"])
    manager = extra.get_grade_manager()
    for sid in students:
        detail = _grade_detail(course_id, sid, knowledge=False)
        records = manager.get_scores(course_id, student_id=sid)
        for row in records or [None]:
            grades.append([sid, detail["display_name"], detail["weighted_total"],
                           row.category if row else None, row.item if row else None,
                           row.score if row else None, row.total if row else None])
    status = workbook.create_sheet("Assignments")
    status.append(["student_id", "assignment_id", "title", "state", "late", "score", "max_score"])
    work, submissions = _assignment_data(course_id)
    for item in work:
        for sid in students:
            row = submissions.get((item["id"], sid), {})
            status.append([sid, item["id"], item["title"], row.get("state", "missing"),
                           bool(row.get("late")), row.get("score"), item["max_score"]])
    activity = workbook.create_sheet("Activity")
    activity.append(["student_id", "sessions", "activity_count", "estimated_hours_last_8_weeks"])
    for sid, rows, messages in _class_activity(course_id):
        activity.append([sid, len(rows), len(_activity_stamps(rows, messages)),
                         round(sum(week["hours"] for week in _weekly(rows, messages)), 4)])
    # User-authored titles/names are literal cells, never Excel formulas.
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    output = io.BytesIO()
    workbook.save(output)
    return Response(output.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition":
                             f'attachment; filename="{quote(course_id, safe="")}-analytics.xlsx"'})


def _cjk_font() -> str | None:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFError, TTFont

    with _FONT_LOCK:
        if "EduInsightCJK" in pdfmetrics.getRegisteredFontNames():
            return "EduInsightCJK"
        candidates = [Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts/kaiu.ttf",
                      Path("/usr/share/fonts/truetype/arphic/ukai.ttc"),
                      Path("/usr/share/fonts/truetype/arphic/uming.ttc")]
        for path in candidates:
            if path.is_file():
                try:
                    pdfmetrics.registerFont(TTFont("EduInsightCJK", str(path)))
                    return "EduInsightCJK"
                except (OSError, ValueError, TTFError):
                    # Some installed CJK files use unsupported outlines. Try the next font.
                    continue
    return None


def _minimal_pdf(lines: list[str]) -> bytes:
    """Dependency-free multipage PDF with ASCII text and a valid xref."""
    pages = [lines[index:index + 45] for index in range(0, len(lines), 45)] or [[]]
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"",
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    page_ids = []
    for page in pages:
        page_id = len(objects) + 1
        page_ids.append(page_id)
        content_id = page_id + 1
        commands = ["BT /F1 10 Tf 45 790 Td 16 TL"]
        for index, line in enumerate(page):
            text = line.encode("ascii", "replace").decode().replace("\\", "\\\\")
            text = text.replace("(", "\\(").replace(")", "\\)").replace("\n", " ")
            commands.append(f"{'T* ' if index else ''}({text}) Tj")
        stream = ("\n".join(commands) + "\nET").encode("ascii")
        objects.extend([
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>".encode(),
            f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream",
        ])
    kids = " ".join(f"{pid} 0 R" for pid in page_ids)
    objects[1] = f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode()
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\n"
                  f"startxref\n{xref}\n%%EOF\n".encode())
    return bytes(output)


def _transcript(course_id: str, student_id: int) -> Response:
    detail = _grade_detail(course_id, student_id, knowledge=False)
    total = detail["weighted_total"]
    lines = ["EduInsight Transcript / 學生成績單", f"Course / 課程: {course_id} {_course_catalog()[course_id]}",
             f"Student / 學生: {student_id} {detail['display_name']}",
             f"Weighted total / 加權總分: {total if total is not None else 'N/A'}"]
    for category in detail["category_scores"]:
        lines.append(f"{category['category']} ({category['weight']}%)")
        lines.extend(f"{item['item']}: {item['score']} / {item['total']}" for item in category["items"])
    work, submissions = _assignment_data(course_id)
    for item in work:
        row = submissions.get((item["id"], student_id), {})
        lines.append(f"{item['id']} {item['title']}: {row.get('state', 'missing')} "
                     f"{row.get('score') if row.get('score') is not None else 'N/A'} / {item['max_score']}")
    try:
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
        font = _cjk_font()
    except ImportError:
        font = None
    if font:
        output = io.BytesIO()
        style = ParagraphStyle("transcript", fontName=font, fontSize=11, leading=18, wordWrap="CJK")
        content = []
        for line in lines:
            content.extend([Paragraph(escape(line), style), Spacer(1, 6)])
        SimpleDocTemplate(output).build(content)
        pdf = output.getvalue()
    else:
        lines.insert(0, "CJK font unavailable: ASCII text and IDs only.")
        pdf = _minimal_pdf(lines)
    headers = {"Content-Disposition":
               f'attachment; filename="{quote(course_id, safe="")}-{student_id}-transcript.pdf"',
               "X-Transcript-Text": "CJK" if font else "ASCII",
               "Cache-Control": "private, no-store"}
    if not font:
        headers["X-Transcript-Limitation"] = "CJK font unavailable; ASCII text and IDs only"
    return Response(pdf, media_type="application/pdf", headers=headers)


@router.get("/class-analytics/{course_id}/student/{moodle_user_id}/transcript.pdf")
def transcript(course_id: str, moodle_user_id: int,
               user: AuthUser = Depends(get_current_user)) -> Response:
    _student_access(user, course_id, moodle_user_id)
    return _transcript(course_id, moodle_user_id)


@router.get("/me/courses/{course_id}/transcript.pdf")
def own_transcript(course_id: str, user: AuthUser = Depends(require_role("student"))) -> Response:
    if user.moodle_user_id is None:
        raise HTTPException(403, "Student Moodle identity is required")
    _student_access(user, course_id, user.moodle_user_id)
    return _transcript(course_id, user.moodle_user_id)
