"""Tests for RAG API endpoints (course material upload, search, chat integration)."""

from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient
from litemem import Memory

from eduinsight import app as app_module
from eduinsight.rag import CourseRAG


@pytest.fixture
def client():
    """Create a test client with in-memory database."""
    mem = Memory(":memory:")
    rag = CourseRAG(mem)

    with TestClient(app_module.app, raise_server_exceptions=False) as c:
        # Save originals created by lifespan, override for test
        orig_llm = app_module._llm
        app_module._memory = mem
        app_module._assistant = app_module.LearningAssistant(mem)
        app_module._rag = rag
        app_module._llm = None
        app_module._quiz = None
        yield c
        # Restore original _llm so lifespan cleanup (__aexit__) works
        app_module._llm = orig_llm


def _make_test_pdf(text: str = "Binary search works by dividing the array in half.") -> bytes:
    """Create a minimal PDF in memory."""
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), text)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def _make_test_pptx(text: str = "Linked list traversal example.") -> bytes:
    """Create a minimal PPTX in memory."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    txBox.text_frame.text = text
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


class TestMaterialUpload:
    def test_upload_pdf(self, client):
        pdf_bytes = _make_test_pdf()
        resp = client.post(
            "/courses/CS101/materials",
            files={"file": ("lecture01.pdf", pdf_bytes, "application/pdf")},
        )
        assert resp.status_code == 200
        data = resp.json()
        assert data["status"] == "indexed"
        assert data["filename"] == "lecture01.pdf"
        assert data["course_id"] == "CS101"
        assert data["chunks_indexed"] >= 1
        assert data["total_pages"] == 1

    def test_upload_pptx(self, client):
        pptx_bytes = _make_test_pptx()
        resp = client.post(
            "/courses/CS101/materials",
            files={"file": ("slides.pptx", pptx_bytes, "application/vnd.openxmlformats")},
        )
        assert resp.status_code == 200
        data = resp.json()
        assert data["status"] == "indexed"
        assert data["filename"] == "slides.pptx"

    def test_upload_unsupported_format(self, client):
        resp = client.post(
            "/courses/CS101/materials",
            files={"file": ("notes.txt", b"hello", "text/plain")},
        )
        assert resp.status_code == 400
        assert "Unsupported" in resp.json()["detail"]


class TestMaterialList:
    def test_list_empty(self, client):
        resp = client.get("/courses/CS101/materials")
        assert resp.status_code == 200
        assert resp.json()["documents"] == []

    def test_list_after_upload(self, client):
        pdf_bytes = _make_test_pdf()
        client.post(
            "/courses/CS101/materials",
            files={"file": ("lecture01.pdf", pdf_bytes, "application/pdf")},
        )
        resp = client.get("/courses/CS101/materials")
        assert resp.json()["documents"] == ["lecture01.pdf"]

    def test_list_multiple(self, client):
        client.post(
            "/courses/CS101/materials",
            files={"file": ("a.pdf", _make_test_pdf("A"), "application/pdf")},
        )
        client.post(
            "/courses/CS101/materials",
            files={"file": ("b.pdf", _make_test_pdf("B"), "application/pdf")},
        )
        resp = client.get("/courses/CS101/materials")
        assert sorted(resp.json()["documents"]) == ["a.pdf", "b.pdf"]

    def test_courses_isolated(self, client):
        client.post(
            "/courses/CS101/materials",
            files={"file": ("a.pdf", _make_test_pdf(), "application/pdf")},
        )
        resp = client.get("/courses/CS202/materials")
        assert resp.json()["documents"] == []


class TestMaterialDelete:
    def test_delete_existing(self, client):
        client.post(
            "/courses/CS101/materials",
            files={"file": ("lecture.pdf", _make_test_pdf(), "application/pdf")},
        )
        resp = client.delete("/courses/CS101/materials/lecture.pdf")
        assert resp.status_code == 200
        data = resp.json()
        assert data["status"] == "removed"
        assert data["chunks_removed"] >= 1
        # Verify it's gone
        assert client.get("/courses/CS101/materials").json()["documents"] == []

    def test_delete_nonexistent(self, client):
        resp = client.delete("/courses/CS101/materials/nope.pdf")
        data = resp.json()
        assert data["status"] == "not_found"
        assert data["chunks_removed"] == 0


class TestChatWithRAG:
    def test_chat_without_course_id(self, client):
        """Chat without course_id should work as before (no RAG)."""
        resp = client.post("/chat", json={
            "moodle_user_id": 1001,
            "message": "What is binary search?",
        })
        assert resp.status_code == 200
        data = resp.json()
        assert "material_context" in data
        assert data["material_context"] == []

    def test_chat_with_course_id_and_materials(self, client):
        """Chat with course_id should include RAG context."""
        # Upload material
        client.post(
            "/courses/CS101/materials",
            files={"file": ("algo.pdf", _make_test_pdf(
                "Binary search is an efficient algorithm that finds a target value "
                "in a sorted array by repeatedly dividing the search interval in half."
            ), "application/pdf")},
        )
        # Chat with course_id
        resp = client.post("/chat", json={
            "moodle_user_id": 1001,
            "message": "What is binary search?",
            "course_id": "CS101",
        })
        assert resp.status_code == 200
        data = resp.json()
        # Should have material references
        assert len(data["material_context"]) > 0
        assert any("algo.pdf" in ref for ref in data["material_context"])

    def test_chat_with_empty_course(self, client):
        """Chat with course_id but no materials should still work."""
        resp = client.post("/chat", json={
            "moodle_user_id": 1001,
            "message": "What is binary search?",
            "course_id": "EMPTY",
        })
        assert resp.status_code == 200
        assert resp.json()["material_context"] == []
