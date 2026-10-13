"""Only the two new routers: real JWTs, isolated databases and real data stores."""

from datetime import UTC, datetime, timedelta
from io import BytesIO

import pytest
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from litemem import Memory
from openpyxl import load_workbook

from eduinsight import achievements as badges
from eduinsight import analytics_extra as extra
from eduinsight import assignments, auth, course_structure, sessions
from eduinsight import class_analytics as ca
from eduinsight.config import settings
from eduinsight.grades import GradeManager

NOW = datetime(2026, 10, 9, 12, tzinfo=extra._TAIPEI)
BASE = "/class-analytics/ds101"
CLASS_PATHS = [f"{BASE}/{suffix}" for suffix in (
    "weekly-hours", "focus-heatmap", "skills", "assignment-completion",
    "golden-hours", "export.xlsx", "grades-detail",
)]
PRIVATE_PATHS = [f"{BASE}/student/1001/transcript.pdf", f"{BASE}/student/1001/feedback",
                 "/me/courses/ds101/transcript.pdf", "/me/achievements", "/me/profile-stats",
                 "/me/calendar"]


@pytest.fixture
async def client(auth_headers, monkeypatch):
    memory = Memory(settings.memory_db_path, embedder=None)
    manager = GradeManager()
    monkeypatch.setattr(extra, "get_memory", lambda: memory)
    monkeypatch.setattr(extra, "get_grade_manager", lambda: manager)
    monkeypatch.setattr(extra, "_now", lambda: NOW)
    app = FastAPI()
    app.include_router(ca.router)
    app.include_router(badges.router)
    app.include_router(assignments.router)
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        yield client


def seed_session(sid="own", student=1001, course="ds101", start=None, end=None):
    start = NOW.timestamp() if start is None else start
    end = start if end is None else end
    with sessions._db() as conn:
        conn.execute("INSERT INTO chat_sessions VALUES (?, ?, ?, 'Test', ?, ?)",
                     (sid, student, course, start, end))
        conn.commit()
    from eduinsight.courses import enroll
    enroll(course, student, source="test")
    return sid


def message(*, sid="own", student=1001, when=None, role="user"):
    store = extra.get_memory()._store
    mid = store.store_message(f"moodle:{student}", "Learning", session_id=sid, role=role)
    conn = store._get_conn()
    conn.execute("UPDATE messages SET created_at = ? WHERE id = ?",
                 (NOW.timestamp() if when is None else when, mid))
    conn.commit()


def seed_work(aid="a", course="ds101", title="課程作業", due=None):
    with assignments._db(write=True) as conn:
        conn.execute("INSERT INTO assignments VALUES (?, ?, ?, '', ?, 100, 1, ?, ?)",
                     (aid, course, title, due or NOW.timestamp(), NOW.timestamp(), NOW.timestamp()))
    return aid


def seed_submission(aid="a", sid=1001, state="graded", score=100, late=False, when=None):
    when = NOW.timestamp() - 60 if when is None else when
    with assignments._db(write=True) as conn:
        conn.execute(
            "INSERT INTO assignment_submissions "
            "(assignment_id, student_id, state, late, submitted_at, score, feedback, graded_at, updated_at) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (aid, sid, state, int(late), when if state != "draft" else None,
             score if state == "graded" else None, "清楚完整" if state == "graded" else None,
             when + 30 if state == "graded" else None, when),
        )


async def get(client, auth_headers, path, username="teacher"):
    response = await client.get(path, headers=auth_headers(username))
    assert response.status_code == 200, response.text
    return response.json()


@pytest.mark.parametrize("path", CLASS_PATHS + PRIVATE_PATHS)
@pytest.mark.parametrize("headers", ({}, {"Authorization": "Bearer invalid"}))
async def test_authentication(client, path, headers):
    assert (await client.get(path, headers=headers)).status_code == 401


@pytest.mark.parametrize("path", CLASS_PATHS)
async def test_class_role_and_course_scope(client, auth_headers, path):
    seed_session()
    assert (await client.get(path, headers=auth_headers("student1001"))).status_code == 403
    assert (await client.get(path, headers=auth_headers("teacher_huang"))).status_code == 403
    assert (await client.get(path, headers=auth_headers("teacher"))).status_code == 200
    assert (await client.get(path, headers=auth_headers())).status_code == 200
    auth.create_user(auth.RegisterRequest(username="unmapped", password="1234",
                                         display_name="Unmapped", role="teacher"))
    assert (await client.get(path, headers=auth_headers("unmapped"))).status_code == 403
    assert (await client.get(path.replace("ds101", "unknown"), headers=auth_headers())).status_code == 404


async def test_honest_empty_class(client, auth_headers):
    weekly = await get(client, auth_headers, CLASS_PATHS[0])
    assert weekly["student_count"] == 0 and weekly["estimated"] is True
    assert weekly["timezone"] == "Asia/Taipei" and len(weekly["weeks"]) == 8
    assert all(row["total_hours"] == row["average_hours"] == 0 for row in weekly["weeks"])
    assert all(datetime.fromisoformat(row["week_start"]).weekday() == 0 for row in weekly["weeks"])
    assert (await get(client, auth_headers, CLASS_PATHS[1]))["counts"] == [[0] * 24 for _ in range(7)]
    assert (await get(client, auth_headers, CLASS_PATHS[2]))["skills"] == []
    completion = await get(client, auth_headers, CLASS_PATHS[3])
    assert completion["assignments"] == [] and completion["overall_completion_rate"] == 0
    assert (await get(client, auth_headers, CLASS_PATHS[4]))["slots"] == []
    assert (await get(client, auth_headers, CLASS_PATHS[6]))["students"] == []
    response = await client.get(CLASS_PATHS[5], headers=auth_headers())
    workbook = load_workbook(BytesIO(response.content))
    assert workbook["Grades"].max_row == workbook["Activity"].max_row == 1
    seed_session()
    detail = (await get(client, auth_headers, CLASS_PATHS[6]))["students"][0]
    assert detail["weighted_total"] is None and detail["grade_trend"] == []
    assert detail["knowledge_mastery"] == detail["category_scores"] == []
    assert (await get(client, auth_headers, PRIVATE_PATHS[1], "student1001"))["feedback"] == []


async def test_real_course_activity_and_golden_slots(client, auth_headers):
    start = NOW.replace(hour=10).timestamp()
    seed_session(start=start, end=start + 900)
    message(when=start)
    message(when=start + 900)
    message(when=start + 600, role="assistant")
    seed_session("second", student=1002, start=start - 86400, end=start - 86400 + 1800)
    seed_session("foreign", course="mk201", start=start, end=start + 1800)
    message(sid="foreign", when=start)
    message(sid="foreign", when=start + 1800)
    seed_session("foreign-student", student=1003, course="mk201", start=start, end=start + 1800)
    message(student=1001, sid="legacy", when=start)
    weekly = await get(client, auth_headers, CLASS_PATHS[0])
    assert weekly["student_count"] == 2
    assert weekly["weeks"][-1]["total_hours"] == 0.75
    assert weekly["weeks"][-1]["average_hours"] == 0.375
    heatmap = await get(client, auth_headers, CLASS_PATHS[1])
    assert sum(map(sum, heatmap["counts"])) == 3
    assert heatmap["counts"][4][10] == 2
    golden = await get(client, auth_headers, CLASS_PATHS[4])
    assert golden["slots"][0] == {"weekday": 4, "weekday_name": "Fri", "hour": 10, "count": 2}
    assert len(golden["slots"]) == 2  # Empty slots never advertised as productive.


async def test_editable_catalog_assignments_control_teacher_analytics_scope(client, auth_headers):
    with sessions._db() as conn:
        conn.execute("INSERT INTO course_catalog (id, name, teacher_username) "
                     "VALUES ('new-course', 'New course', 'teacher')")
        conn.execute("UPDATE course_catalog SET name='Renamed', teacher_username='teacher_huang' "
                     "WHERE id='ds101'")
        conn.commit()
    seed_session(course="new-course")
    path = "/class-analytics/new-course/weekly-hours"
    assert (await client.get(path, headers=auth_headers("teacher"))).status_code == 200
    assert (await client.get(path, headers=auth_headers())).status_code == 200
    assert (await client.get(CLASS_PATHS[0], headers=auth_headers("teacher_huang"))).status_code == 200
    assert (await client.get(CLASS_PATHS[0], headers=auth_headers("teacher"))).status_code == 403
    assert (await client.get("/me/courses/new-course/transcript.pdf",
                             headers=auth_headers("student1001"))).status_code == 200


async def test_skills_anonymous_and_unknown_evidence(client, auth_headers):
    seed_session()
    seed_session("second", student=1002)
    seed_session("foreign", student=1003, course="mk201")
    memory = extra.get_memory()
    memory.add("moodle:1001", "[SQL] mastered SELECT", category="mastered", extract=False)
    memory.add("moodle:1002", "[SQL] struggling SELECT", category="struggling", extract=False)
    memory.add("moodle:1001", "[Loops] syntax explanation", category="general", extract=False)
    memory.add("moodle:1003", "[SECRET] mastered", category="mastered", extract=False)
    result = await get(client, auth_headers, CLASS_PATHS[2])
    radar = {row["topic"]: row for row in result["skills"]}
    assert result["anonymised"] is True and result["knowledge_scope"] == "student"
    assert set(radar) == {"SQL", "Loops"}
    assert radar["SQL"]["score"] == 50
    assert radar["Loops"]["score"] is None
    assert "moodle_user_id" not in str(result) and "display_name" not in str(result)


async def test_assignment_counts_and_course_scoping(client, auth_headers):
    seed_session()
    seed_session("second", student=1002)
    seed_session("foreign", student=1003, course="mk201")
    seed_work()
    seed_work("draft")
    seed_work("foreign", course="mk201", title="SECRET")
    seed_submission()
    seed_submission(sid=1002, state="submitted", late=True)
    seed_submission(sid=9999)  # Non-enrolled submission must be excluded.
    seed_submission(aid="draft", state="draft")
    result = await get(client, auth_headers, CLASS_PATHS[3])
    assert result["overall_completion_rate"] == 50
    assert result["assignments"][0] == {"assignment_id": "a", "title": "課程作業",
                                         "submitted": 2, "graded": 1, "late": 1, "missing": 0}
    assert result["assignments"][1]["missing"] == 2
    assert len(result["assignments"]) == 2


async def test_grades_categories_mastery_and_chronological_trend(client, auth_headers):
    seed_session()
    seed_session("no-scores", student=1002)
    manager = extra.get_grade_manager()
    manager.set_categories("ds101", [{"name": "HW", "weight": 40}, {"name": "Quiz", "weight": 60}])
    first = manager.record_score("ds101", student_id=1001, category="HW", item="Later", score=80)
    first.created_at = NOW.timestamp() - 100
    second = manager.record_score("ds101", student_id=1001, category="Quiz", item="Earlier", score=45, total=50)
    second.created_at = NOW.timestamp() - 200
    manager.record_score("ds101", student_id=9999, category="HW", item="SECRET", score=100)
    manager.set_categories("mk201", [{"name": "Foreign", "weight": 100}])
    manager.record_score("mk201", student_id=1001, category="Foreign", item="SECRET", score=100)
    extra.get_memory().add("moodle:1001", "[SQL] mastered", category="mastered", extract=False)
    result = await get(client, auth_headers, CLASS_PATHS[6])
    assert len(result["students"]) == 2
    own = result["students"][0]
    assert own["display_name"] == auth.get_user_by_username("student1001")["display_name"]
    assert own["weighted_total"] == 86
    assert [row["item"] for row in own["grade_trend"]] == ["Earlier", "Later"]
    assert own["category_scores"][1]["average_pct"] == 90
    assert own["knowledge_mastery"][0]["score"] == 100
    assert result["students"][1]["weighted_total"] is None
    assert "SECRET" not in str(result)
    first.created_at = second.created_at = 0
    own = (await get(client, auth_headers, CLASS_PATHS[6]))["students"][0]
    assert [row["item"] for row in own["grade_trend"]] == ["Later", "Earlier"]
    assert all(row["at"] is None for row in own["grade_trend"])


async def test_xlsx_contains_only_course_roster_grades_status_activity(client, auth_headers):
    seed_session()
    seed_session("foreign", student=1003, course="mk201")
    seed_session("other-course", course="mk201")
    manager = extra.get_grade_manager()
    for course_id in ("ds101", "mk201"):
        manager.set_categories(course_id, [{"name": "HW", "weight": 100}])
    manager.record_score("ds101", student_id=1001, category="HW", item="Own", score=75)
    manager.record_score("ds101", student_id=9999, category="HW", item="SECRET OUTSIDER", score=100)
    manager.record_score("mk201", student_id=1001, category="HW", item="SECRET FOREIGN", score=100)
    seed_work(title="=1+1")
    seed_work("foreign", course="mk201", title="SECRET FOREIGN")
    seed_submission()
    response = await client.get(CLASS_PATHS[5], headers=auth_headers("teacher"))
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert 'filename="ds101-analytics.xlsx"' in response.headers["content-disposition"]
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["Course", "Grades", "Assignments", "Activity"]
    values = [cell.value for sheet in workbook for row in sheet for cell in row]
    assert not any("SECRET" in str(value) for value in values)
    assert 1003 not in values and 9999 not in values and "mk201" not in values
    assert workbook["Assignments"]["C2"].value == "=1+1"
    assert workbook["Assignments"]["C2"].data_type == "s"
    assert workbook["Assignments"]["D2"].value == "graded"
    assert workbook["Grades"]["C2"].value == 75


@pytest.mark.parametrize("suffix", ("transcript.pdf", "feedback"))
async def test_transcript_feedback_ownership(client, auth_headers, suffix):
    seed_session()
    seed_session("other", student=1002)
    path = f"{BASE}/student/1001/{suffix}"
    assert (await client.get(path, headers=auth_headers("student1002"))).status_code == 403
    assert (await client.get(path, headers=auth_headers("teacher_huang"))).status_code == 403
    assert (await client.get(path.replace("1001", "1003"), headers=auth_headers())).status_code == 403
    row_id = auth.get_user_by_username("student1001")["id"]
    assert row_id != 1001
    assert (await client.get(path.replace("1001", str(row_id)), headers=auth_headers("student1001"))).status_code == 403
    for username in ("student1001", "teacher", "admin"):
        assert (await client.get(path, headers=auth_headers(username))).status_code == 200
    if suffix == "transcript.pdf":
        response = await client.get(PRIVATE_PATHS[2], headers=auth_headers("student1001"))
        assert response.status_code == 200 and response.content.startswith(b"%PDF")
        assert len(response.content) > 300
        assert (await client.get(PRIVATE_PATHS[2], headers=auth_headers("teacher"))).status_code == 403
        assert (await client.get(PRIVATE_PATHS[2].replace("ds101", "mk201"),
                                 headers=auth_headers("student1001"))).status_code == 403


async def test_feedback_reads_existing_grading(client, auth_headers, monkeypatch):
    # The parallel catalog implementation currently initializes its catalog in
    # assignment write transactions. Use our read-only course authorization to
    # exercise the actual submission/grading functions without that unrelated lock.
    monkeypatch.setattr(assignments, "_check_course_access", ca._check_course_access)
    seed_session()
    seed_session("foreign", course="mk201")
    due = datetime.now(UTC) + timedelta(days=1)
    seed_work(due=due.timestamp())
    response = await client.put("/assignments/a/submission", headers=auth_headers("student1001"),
                                json={"state": "submitted", "text": "answer"})
    assert response.status_code == 200
    response = await client.post("/assignments/a/submissions/1001/grade", headers=auth_headers("teacher"),
                                 json={"score": 85, "feedback": "請補上推導"})
    assert response.status_code == 200
    seed_work("foreign", course="mk201")
    seed_submission("foreign")
    result = await get(client, auth_headers, PRIVATE_PATHS[1], "student1001")
    assert len(result["feedback"]) == 1
    assert result["feedback"][0]["feedback"] == "請補上推導"
    assert result["feedback"][0]["score"] == 85


async def test_pdf_valid_cjk_and_ascii_fallback(client, auth_headers, monkeypatch):
    import fitz

    seed_session()
    seed_work()
    seed_submission()
    response = await client.get(PRIVATE_PATHS[0], headers=auth_headers("student1001"))
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF") and len(response.content) > 300
    with fitz.open(stream=response.content, filetype="pdf") as document:
        text = "".join(page.get_text() for page in document)
        assert "1001" in text and "ds101" in text
        if response.headers["x-transcript-text"] == "CJK":
            assert "學生成績單" in text and "課程作業" in text
        else:
            assert "CJK font unavailable" in text
    monkeypatch.setattr(ca, "_cjk_font", lambda: None)
    response = await client.get(PRIVATE_PATHS[0], headers=auth_headers("student1001"))
    assert response.headers["x-transcript-text"] == "ASCII"
    assert "ASCII" in response.headers["x-transcript-limitation"]
    with fitz.open(stream=response.content, filetype="pdf") as document:
        assert "1001" in document[0].get_text() and "CJK font unavailable" in document[0].get_text()
    with fitz.open(stream=ca._minimal_pdf([f"item {i}" for i in range(100)]), filetype="pdf") as document:
        assert len(document) == 3 and "item 99" in document[2].get_text()


@pytest.mark.parametrize("path", ("/me/achievements", "/me/profile-stats", "/me/calendar"))
async def test_student_only_me_routes(client, auth_headers, path):
    assert (await client.get(path, headers=auth_headers("teacher"))).status_code == 403
    assert (await client.get(path, headers=auth_headers())).status_code == 403
    auth.create_user(auth.RegisterRequest(username="noidentity", password="1234",
                                         display_name="No ID", role="student"))
    # Register assigns a Moodle identity automatically; explicitly emulate an unlinked account.
    with auth._db() as conn:
        conn.execute("UPDATE users SET moodle_user_id = NULL WHERE username = 'noidentity'")
        conn.commit()
    assert (await client.get(path, headers=auth_headers("noidentity"))).status_code == 403


async def test_empty_achievements_stats_and_calendar(client, auth_headers):
    result = await get(client, auth_headers, "/me/achievements", "student1001")
    assert result["earned"] == [] and len(result["locked"]) == 8
    assert all(row["progress"]["current"] == 0 for row in result["locked"])
    stats = await get(client, auth_headers, "/me/profile-stats", "student1001")
    assert stats["completed_courses"] == stats["badges_count"] == stats["forum_posts"] == 0
    assert stats["estimated_learning_hours"] == 0 and stats["hours_period_weeks"] == 8
    assert (await get(client, auth_headers, "/me/calendar", "student1001"))["events"] == []
    seed_session()  # An empty session is not an AI chat.
    result = await get(client, auth_headers, "/me/achievements", "student1001")
    assert result["earned"] == []


def seed_completion(count=1):
    with course_structure._db(write=True) as conn:
        conn.execute("INSERT INTO course_sections VALUES ('section', 'ds101', 1, 'Week 1')")
        for index in range(2):
            conn.execute("INSERT INTO course_activities VALUES (?, 'section', 'reading', 'Reading', ?, NULL)",
                         (f"activity-{index}", index))
            if index < count:
                conn.execute("INSERT INTO course_activity_completions VALUES (?, 1001, ?)",
                             (f"activity-{index}", NOW.timestamp() - 120))


def seed_forum():
    # Mirror forum.py's auth-ID ownership while keeping the router under test independent.
    first = auth.get_user_by_username("student1001")["id"]
    second = auth.get_user_by_username("student1002")["id"]
    with sessions._db() as conn:
        conn.executescript("""
            CREATE TABLE forum_threads (id TEXT PRIMARY KEY, course_id TEXT NOT NULL);
            CREATE TABLE forum_posts (id TEXT PRIMARY KEY, thread_id TEXT, author_id INTEGER, created_at REAL);
        """)
        conn.execute("INSERT INTO forum_threads VALUES ('thread', 'ds101')")
        conn.execute("INSERT INTO forum_threads VALUES ('foreign', 'mk201')")
        for pid, thread, owner in (("own-post", "thread", first), ("other-post", "thread", second),
                                   ("foreign-post", "foreign", first), ("wrong-id", "thread", 1001)):
            conn.execute("INSERT INTO forum_posts VALUES (?, ?, ?, ?)",
                         (pid, thread, owner, NOW.timestamp() - 60))
        conn.commit()


async def test_all_badges_real_evidence_earned_dates_and_own_only(client, auth_headers):
    seed_session()
    # Seven consecutive Taipei dates including today; first five chats are real user messages.
    for offset in range(6, -1, -1):
        message(when=NOW.timestamp() - offset * 86400 - 60)
    for index in range(3):
        seed_work(str(index))
        seed_submission(str(index), when=NOW.timestamp() - 300 + index * 60)
    seed_completion(count=1)
    seed_forum()
    result = await get(client, auth_headers, "/me/achievements", "student1001")
    assert len(result["earned"]) == 8 and result["locked"] == []
    assert all(row["earned_at"] is not None for row in result["earned"])
    earned = {row["id"]: row for row in result["earned"]}
    assert datetime.fromisoformat(earned["five_ai_chats"]["earned_at"]).timestamp() == NOW.timestamp() - 2 * 86400 - 60
    stats = await get(client, auth_headers, "/me/profile-stats", "student1001")
    assert stats["completed_courses"] == 0 and stats["badges_count"] == 8 and stats["forum_posts"] == 1
    seed_session("other", student=1002)
    other = await get(client, auth_headers, "/me/achievements?moodle_user_id=1001", "student1002")
    assert [row["id"] for row in other["earned"]] == ["first_forum_post"]
    with course_structure._db(write=True) as conn:
        conn.execute("INSERT INTO course_activity_completions VALUES ('activity-1', 1001, ?)",
                     (NOW.timestamp() - 30,))
    stats = await get(client, auth_headers, "/me/profile-stats", "student1001")
    assert stats["completed_courses"] == 1 and stats["completed_course_ids"] == ["ds101"]


async def test_badge_progress_not_late_draft_foreign_or_unrelated_grade(client, auth_headers):
    seed_session()
    seed_work()
    seed_submission(state="submitted", late=True)
    seed_work("draft")
    seed_submission("draft", state="draft")
    seed_work("foreign", course="mk201")
    seed_submission("foreign")
    message(when=NOW.timestamp() - 60)
    message(when=NOW.timestamp() + 60)  # Future activity is not earned evidence.
    message(role="assistant")
    manager = extra.get_grade_manager()
    manager.set_categories("mk201", [{"name": "HW", "weight": 100}])
    manager.record_score("mk201", student_id=1001, category="HW", item="Perfect", score=100)
    result = await get(client, auth_headers, "/me/achievements", "student1001")
    assert {row["id"] for row in result["earned"]} == {"first_submission", "first_ai_chat"}
    locked = {row["id"]: row for row in result["locked"]}
    assert locked["on_time_three"]["progress"]["current"] == 0
    assert locked["five_ai_chats"]["progress"]["current"] == 1
    assert locked["seven_day_streak"]["progress"]["current"] == 1


async def test_gradebook_perfect_badge_unknown_timestamp_and_future_exclusion(client, auth_headers):
    seed_session()
    manager = extra.get_grade_manager()
    manager.set_categories("ds101", [{"name": "HW", "weight": 100}])
    record = manager.record_score("ds101", student_id=1001, category="HW", item="Perfect", score=100)
    record.created_at = NOW.timestamp() + 60
    result = await get(client, auth_headers, "/me/achievements", "student1001")
    assert result["earned"] == []
    record.created_at = 0  # Existing grades can lack historical timestamp metadata.
    result = await get(client, auth_headers, "/me/achievements", "student1001")
    assert [row["id"] for row in result["earned"]] == ["perfect_score"]
    assert result["earned"][0]["earned_at"] is None


async def test_profile_hours_merge_overlaps_between_own_courses(client, auth_headers):
    start = NOW.timestamp() - 1800
    seed_session(start=start, end=start + 900)
    seed_session("overlap", course="mk201", start=start, end=start + 900)
    seed_session("other", student=1002, start=start, end=start + 1800)
    result = await get(client, auth_headers, "/me/profile-stats", "student1001")
    assert result["estimated_learning_hours"] == 0.25


async def test_calendar_crud_merging_ownership_and_validation(client, auth_headers):
    seed_session()
    seed_session("other", student=1002, course="mk201")
    seed_work()
    seed_work("foreign", course="mk201", title="SECRET")
    path = "/me/calendar/events"
    body = {"title": "  複習  ", "start": "2026-10-10T14:00:00+08:00",
            "end": "2026-10-10T15:00:00+08:00", "notes": "準備考試"}
    assert (await client.post(path, json=body)).status_code == 401
    assert (await client.post(path, json=body, headers=auth_headers("teacher"))).status_code == 403
    response = await client.post(path, json=body, headers=auth_headers("student1001"))
    assert response.status_code == 201
    event = response.json()
    assert event["title"] == "複習" and event["source"] == "custom"
    event_path = f"{path}/{event['id']}"
    merged = await get(client, auth_headers, "/me/calendar", "student1001")
    assert [row["source"] for row in merged["events"]] == ["assignment", "custom"]
    assert "SECRET" not in str(merged)
    other = await get(client, auth_headers, "/me/calendar", "student1002")
    assert [row["id"] for row in other["events"]] == ["assignment:foreign"]
    assert (await client.delete(event_path)).status_code == 401
    assert (await client.put(event_path, json=body, headers=auth_headers("student1002"))).status_code == 404
    assert (await client.delete(event_path, headers=auth_headers("student1002"))).status_code == 404
    updated = await client.put(event_path, json=body | {"title": "更新"}, headers=auth_headers("student1001"))
    assert updated.status_code == 200 and updated.json()["title"] == "更新"
    assert (await client.delete(event_path, headers=auth_headers("student1001"))).status_code == 204
    assert (await client.delete(event_path, headers=auth_headers("student1001"))).status_code == 404
    assert len((await get(client, auth_headers, "/me/calendar", "student1001"))["events"]) == 1
    for invalid in (body | {"title": " "}, body | {"start": "2026-10-10T14:00:00"},
                    body | {"end": "2026-10-09T14:00:00+08:00"}, body | {"student_id": 1002}):
        assert (await client.post(path, json=invalid, headers=auth_headers("student1001"))).status_code == 422
    assert (await client.delete(f"{path}/assignment:a", headers=auth_headers("student1001"))).status_code == 404
    response = await client.post(path, json={"title": "只填起點", "start": body["start"]},
                                 headers=auth_headers("student1001"))
    assert response.status_code == 201
    assert response.json()["end"] is None and response.json()["notes"] == ""
